import os
import json
import re
import tempfile

import streamlit as st
from pypdf import PdfReader
from docx import Document

from langchain_core.documents import Document as LCDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_ollama import ChatOllama, OllamaEmbeddings
from langchain_chroma import Chroma


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="AI Exam Generator",
    page_icon="🎓",
    layout="wide"
)


# ============================================================
# APPLICATION CONFIGURATION
# ============================================================

LLM_MODEL = "llama3.2"
EMBEDDING_MODEL = "nomic-embed-text"
VECTOR_DATABASE_PATH = "./vectorstore"


# ============================================================
# LOAD LLM
# ============================================================

@st.cache_resource
def get_llm():
    return ChatOllama(
        model=LLM_MODEL,
        temperature=0.3
    )


# ============================================================
# LOAD EMBEDDING MODEL
# ============================================================

@st.cache_resource
def get_embeddings():
    return OllamaEmbeddings(
        model=EMBEDDING_MODEL
    )


# ============================================================
# INITIALIZE MODELS
# ============================================================

llm = get_llm()
embeddings = get_embeddings()


# ============================================================
# SESSION STATE
# ============================================================

if "vectorstore" not in st.session_state:
    st.session_state.vectorstore = None

if "questions" not in st.session_state:
    st.session_state.questions = []

if "answers" not in st.session_state:
    st.session_state.answers = {}

if "exam_generated" not in st.session_state:
    st.session_state.exam_generated = False


# ============================================================
# APPLICATION TITLE
# ============================================================

st.title("🎓 AI Exam Generator")

st.write(
    "Generate examination questions from your study material "
    "using Retrieval-Augmented Generation (RAG) and LLMs."
)

st.divider()


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.header("⚙️ Exam Settings")

number_of_questions = st.sidebar.slider(
    "Number of Questions",
    min_value=5,
    max_value=20,
    value=10
)

difficulty = st.sidebar.selectbox(
    "Difficulty Level",
    [
        "Easy",
        "Medium",
        "Hard",
        "Mixed"
    ]
)

st.sidebar.info(
    "Upload your study material and generate "
    "an AI-powered examination."
)


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(uploaded_file):

    reader = PdfReader(uploaded_file)

    complete_text = ""

    for page_number, page in enumerate(reader.pages):

        text = page.extract_text()

        if text:

            complete_text += (
                f"\n--- Page {page_number + 1} ---\n"
            )

            complete_text += text

    return complete_text


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(uploaded_file):

    temporary_path = None

    try:

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=".docx"
        ) as temp_file:

            temp_file.write(
                uploaded_file.getvalue()
            )

            temporary_path = temp_file.name

        document = Document(temporary_path)

        text = ""

        for paragraph in document.paragraphs:

            paragraph_text = paragraph.text.strip()

            if paragraph_text:

                text += paragraph_text + "\n"

        return text

    finally:

        if temporary_path and os.path.exists(temporary_path):

            os.remove(temporary_path)


# ============================================================
# TXT TEXT EXTRACTION
# ============================================================

def extract_txt_text(uploaded_file):

    return uploaded_file.getvalue().decode(
        "utf-8",
        errors="ignore"
    )


# ============================================================
# PROCESS UPLOADED FILE
# ============================================================

def process_file(uploaded_file):

    file_name = uploaded_file.name.lower()

    try:

        if file_name.endswith(".pdf"):

            text = extract_pdf_text(uploaded_file)

        elif file_name.endswith(".docx"):

            text = extract_docx_text(uploaded_file)

        elif file_name.endswith(".txt"):

            text = extract_txt_text(uploaded_file)

        else:

            return None

        if not text.strip():

            return None

        document = LCDocument(
            page_content=text,
            metadata={
                "source": uploaded_file.name
            }
        )

        return document

    except Exception as error:

        st.error(
            f"Error reading {uploaded_file.name}: {error}"
        )

        return None


# ============================================================
# CREATE TEXT CHUNKS
# ============================================================

def create_chunks(documents):

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150
    )

    chunks = splitter.split_documents(
        documents
    )

    return chunks


# ============================================================
# CREATE CHROMA VECTOR DATABASE
# ============================================================

def create_vector_database(chunks):

    vectorstore = Chroma(
        collection_name="exam_documents",
        embedding_function=embeddings,
        persist_directory=VECTOR_DATABASE_PATH
    )

    vectorstore.add_documents(chunks)

    return vectorstore


# ============================================================
# STEP 1 - UPLOAD STUDY MATERIAL
# ============================================================

st.header("📚 Step 1: Upload Study Material")

uploaded_files = st.file_uploader(
    "Upload PDF, DOCX or TXT files",
    type=[
        "pdf",
        "docx",
        "txt"
    ],
    accept_multiple_files=True
)


# ============================================================
# PROCESS DOCUMENTS
# ============================================================

if uploaded_files:

    st.write(
        f"📄 {len(uploaded_files)} file(s) selected."
    )

    if st.button(
        "🔍 Process Documents",
        use_container_width=True
    ):

        documents = []

        # ----------------------------------------------------
        # READ FILES
        # ----------------------------------------------------

        with st.spinner(
            "Reading study material..."
        ):

            for uploaded_file in uploaded_files:

                document = process_file(
                    uploaded_file
                )

                if document:

                    documents.append(document)

        # ----------------------------------------------------
        # CHECK DOCUMENTS
        # ----------------------------------------------------

        if not documents:

            st.error(
                "No readable documents were found."
            )

        else:

            st.success(
                f"Successfully read "
                f"{len(documents)} document(s)."
            )

            # ------------------------------------------------
            # CREATE CHUNKS
            # ------------------------------------------------

            with st.spinner(
                "Splitting documents into chunks..."
            ):

                chunks = create_chunks(
                    documents
                )

            st.info(
                f"Created {len(chunks)} text chunks."
            )

            # ------------------------------------------------
            # CREATE VECTOR DATABASE
            # ------------------------------------------------

            with st.spinner(
                "Creating embeddings and vector database..."
            ):

                try:

                    vectorstore = create_vector_database(
                        chunks
                    )

                    st.session_state.vectorstore = vectorstore

                    st.success(
                        "✅ RAG knowledge base created successfully!"
                    )

                except Exception as error:

                    st.error(
                        f"Error creating vector database: {error}"
                    )


# ============================================================
# STEP 2 - GENERATE EXAM
# ============================================================

st.divider()

st.header("🤖 Step 2: Generate Exam")


if st.session_state.vectorstore is None:

    st.warning(
        "Please upload and process study material first."
    )

else:

    if st.button(
        "✨ Generate Exam",
        use_container_width=True
    ):

        with st.spinner(
            "Retrieving study material and generating questions..."
        ):

            try:

                # ------------------------------------------------
                # CREATE RETRIEVER
                # ------------------------------------------------

                retriever = (
                    st.session_state.vectorstore
                    .as_retriever(
                        search_kwargs={
                            "k": 8
                        }
                    )
                )

                # ------------------------------------------------
                # RETRIEVE RELEVANT DOCUMENTS
                # ------------------------------------------------

                retrieved_documents = retriever.invoke(
                    "Important concepts definitions "
                    "principles examples formulas "
                    "and key topics from the study material"
                )

                # ------------------------------------------------
                # CREATE CONTEXT
                # ------------------------------------------------

                context_parts = []

                for document in retrieved_documents:

                    context_parts.append(
                        document.page_content
                    )

                context = "\n\n".join(
                    context_parts
                )

                # ------------------------------------------------
                # LLM PROMPT
                # ------------------------------------------------

                prompt = f"""
You are an expert university examination question generator.

Generate exactly {number_of_questions} multiple-choice questions.

Difficulty level: {difficulty}

Use ONLY the academic material provided below.

IMPORTANT RULES:

1. Do not use outside knowledge.

2. Every question must be answerable
   from the provided academic material.

3. Generate exactly four options
   for every question.

4. Only one option must be correct.

5. Do not create duplicate questions.

6. Questions must be clear and educational.

7. Questions should be suitable for university students.

8. Return ONLY valid JSON.

9. Do not use Markdown.

10. Do not add explanations outside the JSON.

Use exactly this JSON structure:

[
  {{
    "question": "Question text",
    "options": [
      "Option A",
      "Option B",
      "Option C",
      "Option D"
    ],
    "answer": "Option A"
  }}
]

ACADEMIC MATERIAL:

{context}
"""

                # ------------------------------------------------
                # CALL OLLAMA LLM
                # ------------------------------------------------

                response = llm.invoke(prompt)

                response_text = response.content

                # ------------------------------------------------
                # EXTRACT JSON FROM RESPONSE
                # ------------------------------------------------

                json_match = re.search(
                    r"\[.*\]",
                    response_text,
                    re.DOTALL
                )

                if not json_match:

                    st.error(
                        "The AI did not return valid JSON."
                    )

                    st.code(
                        response_text
                    )

                else:

                    json_text = json_match.group()

                    questions = json.loads(
                        json_text
                    )

                    # --------------------------------------------
                    # VALIDATE QUESTIONS
                    # --------------------------------------------

                    valid_questions = []

                    for question in questions:

                        if not isinstance(
                            question,
                            dict
                        ):

                            continue

                        if (
                            "question" not in question
                            or
                            "options" not in question
                            or
                            "answer" not in question
                        ):

                            continue

                        options = question["options"]

                        answer = question["answer"]

                        if not isinstance(
                            options,
                            list
                        ):

                            continue

                        if len(options) != 4:

                            continue

                        if answer not in options:

                            continue

                        valid_questions.append(
                            question
                        )

                    # --------------------------------------------
                    # SAVE QUESTIONS
                    # --------------------------------------------

                    if not valid_questions:

                        st.error(
                            "The AI generated no valid questions."
                        )

                    else:

                        st.session_state.questions = (
                            valid_questions
                        )

                        st.session_state.answers = {}

                        st.session_state.exam_generated = True

                        st.success(
                            f"🎉 Generated "
                            f"{len(valid_questions)} questions!"
                        )

            except json.JSONDecodeError:

                st.error(
                    "The AI returned invalid JSON. "
                    "Please try generating the exam again."
                )

            except Exception as error:

                st.error(
                    f"Error while generating exam: {error}"
                )


# ============================================================
# STEP 3 - TAKE EXAM
# ============================================================

if (
    st.session_state.exam_generated
    and
    st.session_state.questions
):

    st.divider()

    st.header("📝 Step 3: Take Your Exam")

    questions = st.session_state.questions

    # --------------------------------------------------------
    # DISPLAY QUESTIONS
    # --------------------------------------------------------

    for index, question in enumerate(
        questions
    ):

        st.subheader(
            f"Question {index + 1}"
        )

        st.write(
            question["question"]
        )

        answer = st.radio(
            "Select your answer:",
            question["options"],
            key=f"question_{index}",
            index=None
        )

        st.session_state.answers[index] = answer

    st.divider()

    # --------------------------------------------------------
    # SUBMIT EXAM
    # --------------------------------------------------------

    if st.button(
        "✅ Submit Exam",
        use_container_width=True
    ):

        score = 0

        unanswered = 0

        # ----------------------------------------------------
        # CALCULATE SCORE
        # ----------------------------------------------------

        for index, question in enumerate(
            questions
        ):

            user_answer = (
                st.session_state.answers.get(
                    index
                )
            )

            correct_answer = (
                question["answer"]
            )

            if user_answer is None:

                unanswered += 1

            elif user_answer == correct_answer:

                score += 1

        # ----------------------------------------------------
        # CALCULATE RESULT
        # ----------------------------------------------------

        total_questions = len(
            questions
        )

        percentage = (
            score / total_questions
        ) * 100

        # ----------------------------------------------------
        # DISPLAY RESULT
        # ----------------------------------------------------

        st.divider()

        st.header("📊 Exam Result")

        col1, col2, col3 = st.columns(3)

        with col1:

            st.metric(
                "Score",
                f"{score}/{total_questions}"
            )

        with col2:

            st.metric(
                "Percentage",
                f"{percentage:.1f}%"
            )

        with col3:

            if percentage >= 80:

                performance = "Excellent"

            elif percentage >= 60:

                performance = "Good"

            elif percentage >= 40:

                performance = "Average"

            else:

                performance = "Needs Improvement"

            st.metric(
                "Performance",
                performance
            )

        # ----------------------------------------------------
        # UNANSWERED QUESTIONS
        # ----------------------------------------------------

        if unanswered > 0:

            st.warning(
                f"You left {unanswered} "
                f"question(s) unanswered."
            )

        # ----------------------------------------------------
        # ANSWER REVIEW
        # ----------------------------------------------------

        st.divider()

        st.subheader(
            "📖 Answer Review"
        )

        for index, question in enumerate(
            questions
        ):

            user_answer = (
                st.session_state.answers.get(
                    index
                )
            )

            correct_answer = (
                question["answer"]
            )

            st.write(
                f"**Q{index + 1}. "
                f"{question['question']}**"
            )

            if user_answer == correct_answer:

                st.success(
                    f"Your answer: "
                    f"{user_answer} ✓"
                )

            else:

                if user_answer:

                    st.error(
                        f"Your answer: "
                        f"{user_answer}"
                    )

                else:

                    st.warning(
                        "Not answered"
                    )

                st.info(
                    f"Correct answer: "
                    f"{correct_answer}"
                )